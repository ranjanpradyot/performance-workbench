"""Record Q3-2026 learning evidence and mark the Learning KPI Yes.

Santosh and Krishna use the Bitbucket paths supplied for the review.
Utkarsh uses the same repository and quarter folder: 2026/Utkarsh/Q3.
"""

from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[1]
REPO = "https://bitbucket.org/pranjan/performance-learning-development/src/main"

PEOPLE = {
    "Utkarsh": {
        "file": "Utkarsh_OKR_Checklist_Linked_0.4.xlsx",
        "kpi": "Learning & Development",
        "link": f"{REPO}/2026/Utkarsh/Q3/",
        "actual": "Python 3.12 suite: PayUp expense splitter, Word Scrambler, and Karaoke Queue. 17 unit tests passed. Updated POC to be shared by 23 Oct 2026.",
        "row": [
            "Q3-2026",
            "Learning & Development",
            "Python 3.12 core suite — PayUp, Word Scrambler, Karaoke Queue",
            "Three standalone applications covering arithmetic and penny-balancing, loops and anagram checks, and defensive queue handling.",
            "100% traceability across 62 Python core lessons. 17 automated unit tests passed. Sub-millisecond processing on all three modules.",
            "Python 3.12+, standard-library unittest, CLI",
            f"{REPO}/2026/Utkarsh/Q3/",
            "Yes",
            "PayUp Expense Splitter, Word Scrambler Game, and Karaoke Queue Manager",
            "Q3 review on 5 Oct 2026 and POC demo on 15 Oct 2026. Updated code to be shared by 23 Oct 2026.",
        ],
    },
    "Krishna": {
        "file": "Krishna_OKR_Checklist_Linked_0.4.xlsx",
        "kpi": "Learning & Development",
        "link": f"{REPO}/2026/Krishnaraj/Q3/",
        "actual": "Optimized EPC.vwProductStructure from about 26s to 0.6s and CPU from about 181s to 4.3s, with 2,091,288 rows unchanged. Follow-up on execution plans by 29 Oct 2026.",
        "row": [
            "Q3-2026",
            "Learning & Development",
            "SQL query optimization — EPC.vwProductStructure",
            "A slow SQL Server view in SDHOperational. Baseline metrics, execution plans, and STATISTICS IO/TIME showed cost from repeated references to a complex view.",
            "Execution about 26s to 0.6s (about 97.7%). CPU about 181s to 4.3s (about 97.6%). Row count 2,091,288 matched, with selected records checked.",
            "SQL Server, execution plans, STATISTICS IO, STATISTICS TIME",
            f"{REPO}/2026/Krishnaraj/Q3/",
            "Yes",
            "Restructured query with before-and-after performance metrics",
            "Q3 review summary. Execution-plan follow-up to be shared by 29 Oct 2026.",
        ],
    },
    "Santosh": {
        "file": "Santosh_OKR_Checklist_Linked_0.4.xlsx",
        "kpi": "L&D",
        "link": f"{REPO}/2026/Santosh/Q3/",
        "actual": "Python automation POC for QAD purchase orders: validate open POs, classify discrete and schedule orders, and write a supplier summary. Demonstrated 15 Oct 2026.",
        "row": [
            "Q3-2026",
            "L&D",
            "Intermediate Python and QAD purchase-order automation. Reads PO data, validates open transactions, classifies discrete and schedule POs, calculates value, and writes a supplier CSV summary. Learning demonstration, not a production QAD integration. Q3 review 5 Oct 2026; POC demo 15 Oct 2026.",
            f"{REPO}/2026/Santosh/Q3/",
            "Yes",
        ],
    },
}


def header_column(sheet, name: str) -> int:
    for cell in sheet[1]:
        if cell.value == name:
            return cell.column
    raise SystemExit(f"{sheet.title} has no {name} column")


def learning_row(sheet, quarter: str):
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row, 1).value == quarter:
            return row
    return None


def write_link(cell, url: str, style_from) -> None:
    cell.value = url
    cell.hyperlink = url
    cell.font = Font(name="Calibri", color="FF0563C1", underline="single", size=style_from.font.size or 11)
    cell.alignment = Alignment(wrap_text=True, vertical="top")


def upsert_learning(sheet, values: list, link_column: int) -> None:
    found = learning_row(sheet, "Q3-2026")
    if found:
        target = found
    else:
        after = learning_row(sheet, "Q2-2026")
        if not after:
            raise SystemExit(f"{sheet.title} has no Q2-2026 row to insert after")
        sheet.insert_rows(after + 1)
        target = after + 1
        for column in range(1, len(values) + 1):
            source = sheet.cell(after, column)
            dest = sheet.cell(target, column)
            if source.has_style:
                dest.font = copy(source.font)
                dest.border = copy(source.border)
                dest.fill = copy(source.fill)
                dest.alignment = copy(source.alignment)
                dest.number_format = source.number_format
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(target, column)
        if column == link_column:
            write_link(cell, value, sheet.cell(target, 1))
        else:
            cell.value = value
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def main() -> None:
    for person, spec in PEOPLE.items():
        path = ROOT / "fixtures" / "workbooks" / spec["file"]
        book = load_workbook(path)
        learning = book["Learning_Evidence"]
        link_column = 4 if person == "Santosh" else 7
        upsert_learning(learning, spec["row"], link_column)
        kpi = book["KPI_2026"]
        actual_col = header_column(kpi, "Q3Actual")
        achieved_col = header_column(kpi, "Q3Achieved")
        kpi_row = next(row for row in range(2, kpi.max_row + 1) if kpi.cell(row, 1).value == spec["kpi"])
        kpi.cell(kpi_row, actual_col).value = f"{spec['actual']} {spec['link']}"
        kpi.cell(kpi_row, actual_col).alignment = Alignment(wrap_text=True, vertical="top")
        kpi.cell(kpi_row, achieved_col).value = "Yes"
        document = book["Document_Summary"]
        note = document.max_row + 1
        document.cell(note, 1).value = "0.4"
        document.cell(note, 2).value = "Q3 learning"
        document.cell(note, 3).value = f"Q3-2026 Learning is Yes. Evidence: {spec['link']}"
        document.cell(note, 3).alignment = Alignment(wrap_text=True)
        book.save(path)
        print(person, "Q3 learning Yes", spec["link"])


if __name__ == "__main__":
    main()
